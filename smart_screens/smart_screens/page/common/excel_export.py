"""
Excel export utilities for aggregated reports
Handles Excel generation with formatting for both SPP and BatCom reports
"""
import frappe
from io import BytesIO
from datetime import datetime

def export_aggregated_data_to_excel(data, grand_total, report_title, filters, warehouse_filter='', item_groups=None):
	"""
Export aggregated report data to Excel with formatting

Args:
data: Aggregated data list
grand_total: Grand total dictionary
report_title: Title of the report
filters: Report filters (from_date, to_date, etc.)
warehouse_filter: Selected warehouse filter (optional)
item_groups: List of item groups to include in columns (e.g., ['Mat', 'Products', 'Finished Product'])

Returns:
dict: {success: bool, file_url: str, file_name: str} or {success: False, error: str}
"""
	import openpyxl
	from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
	from openpyxl.utils import get_column_letter
	
	try:
		if not data:
			return {
				"success": False,
				"error": "No data available for export"
			}
		
		# If item_groups not specified, try to detect from grand_total keys
		if not item_groups:
			item_groups = [k for k in grand_total.keys() if k != "Total"]
		
		# Create Excel workbook
		wb = openpyxl.Workbook()
		ws = wb.active
		ws.title = report_title[:31]  # Excel sheet name limit
		
		# Styling
		header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
		header_font = Font(bold=True, color="FFFFFF", size=11)
		total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
		total_font = Font(bold=True, size=10)
		border = Border(
left=Side(style='thin'),
right=Side(style='thin'),
top=Side(style='thin'),
bottom=Side(style='thin')
)
		center_align = Alignment(horizontal='center', vertical='center')
		right_align = Alignment(horizontal='right', vertical='center')
		
		# Report Title
		current_row = 1
		total_cols = 1 + (len(item_groups) * 4) + 4  # Common Code + (each group * 4 cols) + Total cols
		ws.merge_cells(f'A{current_row}:{get_column_letter(total_cols)}{current_row}')
		title_cell = ws[f'A{current_row}']
		title_cell.value = f"{report_title} - Common Code Summary"
		title_cell.font = Font(bold=True, size=14, color="366092")
		title_cell.alignment = center_align
		current_row += 1
		
		# Filter information
		ws.merge_cells(f'A{current_row}:{get_column_letter(total_cols)}{current_row}')
		filter_info = ws[f'A{current_row}']
		filter_text = f"Period: {filters.get('from_date', '')} to {filters.get('to_date', '')}"
		if warehouse_filter and warehouse_filter != '':
			filter_text += f" | Warehouse: {warehouse_filter}"
		else:
			filter_text += " | All Warehouses"
		filter_info.value = filter_text
		filter_info.font = Font(italic=True, size=10)
		filter_info.alignment = center_align
		current_row += 2
		
		# Headers
		headers = ["Common Code"]
		for item_group in item_groups:
			headers.extend([
f"{item_group} Opening",
f"{item_group} In",
f"{item_group} Out",
f"{item_group} Balance"
])
		headers.extend(["Total Opening", "Total In", "Total Out", "Total Balance"])
		
		for col_num, header in enumerate(headers, 1):
			cell = ws.cell(row=current_row, column=col_num)
			cell.value = header
			cell.fill = header_fill
			cell.font = header_font
			cell.alignment = center_align
			cell.border = border
		
		current_row += 1
		
		# Data rows
		for row_data in data:
			col_num = 1
			
			# Common Code
			cell = ws.cell(row=current_row, column=col_num)
			cell.value = row_data.get('common_code', '')
			cell.alignment = center_align
			cell.border = border
			col_num += 1
			
			# Item group columns
			for item_group in item_groups:
				group_data = row_data.get(item_group, {})
				for key in ['opening_qty', 'in_qty', 'out_qty', 'balance_qty']:
					cell = ws.cell(row=current_row, column=col_num)
					cell.value = group_data.get(key, 0)
					cell.alignment = right_align
					cell.border = border
					cell.number_format = '#,##0.00'
					col_num += 1
			
			# Total columns
			total = row_data.get('Total', {})
			for key in ['opening_qty', 'in_qty', 'out_qty', 'balance_qty']:
				cell = ws.cell(row=current_row, column=col_num)
				cell.value = total.get(key, 0)
				cell.alignment = right_align
				cell.border = border
				cell.number_format = '#,##0.00'
				cell.font = Font(bold=True)
				col_num += 1
			
			current_row += 1
		
		# Grand Total Row
		col_num = 1
		cell = ws.cell(row=current_row, column=col_num)
		cell.value = "Grand Total"
		cell.fill = total_fill
		cell.font = total_font
		cell.alignment = center_align
		cell.border = border
		col_num += 1
		
		# Add grand total values
		for item_group in item_groups:
			group_data = grand_total.get(item_group, {})
			for key in ['opening_qty', 'in_qty', 'out_qty', 'balance_qty']:
				cell = ws.cell(row=current_row, column=col_num)
				cell.value = group_data.get(key, 0)
				cell.fill = total_fill
				cell.font = total_font
				cell.alignment = right_align
				cell.border = border
				cell.number_format = '#,##0.00'
				col_num += 1
		
		# Total grand total
		total_data = grand_total.get('Total', {})
		for key in ['opening_qty', 'in_qty', 'out_qty', 'balance_qty']:
			cell = ws.cell(row=current_row, column=col_num)
			cell.value = total_data.get(key, 0)
			cell.fill = total_fill
			cell.font = total_font
			cell.alignment = right_align
			cell.border = border
			cell.number_format = '#,##0.00'
			col_num += 1
		
		# Adjust column widths
		ws.column_dimensions['A'].width = 15  # Common Code
		for col in range(2, total_cols + 1):
			ws.column_dimensions[get_column_letter(col)].width = 12
		
		# Freeze panes (freeze header row)
		ws.freeze_panes = ws['A5']
		
		# Save to BytesIO
		file_data = BytesIO()
		wb.save(file_data)
		file_data.seek(0)
		
		# Generate filename
		timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
		safe_warehouse = warehouse_filter.replace(" ", "_").replace("-", "_") if warehouse_filter else "All_Warehouses"
		safe_title = report_title.replace(" ", "_")
		filename = f"{safe_title}_{safe_warehouse}_{timestamp}.xlsx"
		
		# Save file to Frappe File Manager
		file_doc = frappe.get_doc({
"doctype": "File",
"file_name": filename,
"is_private": 0,
"content": file_data.getvalue()
		})
		file_doc.save(ignore_permissions=True)
		
		return {
			"success": True,
			"file_url": file_doc.file_url,
			"file_name": filename,
			"message": f"Excel file generated successfully with {len(data)} records"
		}
		
	except Exception as e:
		frappe.log_error(
f"Error exporting to Excel: {str(e)}\n{frappe.get_traceback()}",
f"{report_title} - Export Error"
)
		return {
			"success": False,
			"error": str(e)
		}

@frappe.whitelist()
def export_batch_details_to_excel(common_code, batches, filters, warehouse_filter='', item_groups=None):
	"""
Export batch details for a specific common code to Excel

Args:
common_code: The common code to export
batches: Batch data grouped by item group
filters: Report filters
warehouse_filter: Selected warehouse filter
item_groups: List of item groups (e.g., ['Mat', 'Products', 'Finished Product'])

Returns:
dict: {success: bool, file_url: str, file_name: str} or {success: False, error: str}
"""
	import openpyxl
	from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
	from openpyxl.utils import get_column_letter
	
	import json
	
	try:
		# Handle stringified arguments from frappe.call
		if isinstance(batches, str):
			batches = json.loads(batches)
		if isinstance(filters, str):
			filters = json.loads(filters)
		if isinstance(item_groups, str):
			item_groups = json.loads(item_groups)
			
		# If item_groups not specified, get from batches keys
		if not item_groups:
			item_groups = list(batches.keys())
		
		# Create Excel workbook
		wb = openpyxl.Workbook()
		ws = wb.active
		ws.title = f"Code {common_code}"[:31]
		
		# Styling
		header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
		header_font = Font(bold=True, color="FFFFFF", size=11)
		border = Border(
left=Side(style='thin'),
right=Side(style='thin'),
top=Side(style='thin'),
bottom=Side(style='thin')
)
		center_align = Alignment(horizontal='center', vertical='center')
		right_align = Alignment(horizontal='right', vertical='center')
		
		# Title
		current_row = 1
		ws.merge_cells(f'A{current_row}:J{current_row}')
		title_cell = ws[f'A{current_row}']
		title_cell.value = f"Batch Details for Common Code: {common_code}"
		title_cell.font = Font(bold=True, size=14, color="366092")
		title_cell.alignment = center_align
		current_row += 2
		
		# Headers
		headers = [
			"Item Code", "Item Group", "Batch No", "Warehouse",
			"Opening Qty", "In Qty", "Out Qty", "Balance Qty",
			"UOM", "Value"
		]
		
		for col_num, header in enumerate(headers, 1):
			cell = ws.cell(row=current_row, column=col_num)
			cell.value = header
			cell.fill = header_fill
			cell.font = header_font
			cell.alignment = center_align
			cell.border = border
		
		current_row += 1
		
		# Add data for each item group
		for item_group in item_groups:
			batch_list = batches.get(item_group, [])
			
			if not batch_list:
				continue
			
			# Group header
			ws.merge_cells(f'A{current_row}:J{current_row}')
			group_cell = ws[f'A{current_row}']
			group_cell.value = f"{item_group} ({len(batch_list)} batches)"
			group_cell.font = Font(bold=True, size=11)
			group_cell.fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
			current_row += 1
			
			# Batch rows
			for batch in batch_list:
				col_num = 1
				
				# Item Code
				cell = ws.cell(row=current_row, column=col_num)
				cell.value = batch.get('item', '')
				cell.border = border
				col_num += 1
				
				# Item Group
				cell = ws.cell(row=current_row, column=col_num)
				cell.value = batch.get('item_group', '')
				cell.border = border
				col_num += 1
				
				# Batch No
				cell = ws.cell(row=current_row, column=col_num)
				cell.value = batch.get('batch', '')
				cell.border = border
				col_num += 1
				
				# Warehouse
				cell = ws.cell(row=current_row, column=col_num)
				cell.value = batch.get('warehouse', '')
				cell.border = border
				col_num += 1
				
				# Quantities
				for key in ['opening_qty', 'in_qty', 'out_qty', 'balance_qty']:
					cell = ws.cell(row=current_row, column=col_num)
					cell.value = batch.get(key, 0)
					cell.alignment = right_align
					cell.border = border
					cell.number_format = '#,##0.00'
					col_num += 1
				
				# UOM
				cell = ws.cell(row=current_row, column=col_num)
				cell.value = batch.get('uom', 'Nos')
				cell.alignment = center_align
				cell.border = border
				col_num += 1
				
				# Value
				cell = ws.cell(row=current_row, column=col_num)
				cell.value = batch.get('balance_value', 0)
				cell.alignment = right_align
				cell.border = border
				cell.number_format = '#,##0.00'
				col_num += 1
				
				current_row += 1
			
			current_row += 1  # Add spacing between groups
		
		# Adjust column widths
		ws.column_dimensions['A'].width = 18  # Item Code
		ws.column_dimensions['B'].width = 18  # Item Group
		ws.column_dimensions['C'].width = 15  # Batch No
		ws.column_dimensions['D'].width = 25  # Warehouse
		for col in range(5, 11):
			ws.column_dimensions[get_column_letter(col)].width = 12
		
		# Freeze panes
		ws.freeze_panes = ws['A4']
		
		# Save to BytesIO
		file_data = BytesIO()
		wb.save(file_data)
		file_data.seek(0)
		
		# Generate filename
		timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
		filename = f"Batch_Details_{common_code}_{timestamp}.xlsx"
		
		# Save file
		file_doc = frappe.get_doc({
"doctype": "File",
"file_name": filename,
"is_private": 0,
"content": file_data.getvalue()
		})
		file_doc.save(ignore_permissions=True)
		
		return {
			"success": True,
			"file_url": file_doc.file_url,
			"file_name": filename
		}
		
	except Exception as e:
		frappe.log_error(
f"Error exporting batch details: {str(e)}\n{frappe.get_traceback()}",
"Batch Details Export Error"
)
		return {
			"success": False,
			"error": str(e)
		}
def export_rm_data_to_excel(data, grand_total, report_title, filters, warehouse_filter=''):
	"""
	Export Raw Material aggregated report data to Excel with RM-specific columns
	"""
	import openpyxl
	from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
	from openpyxl.utils import get_column_letter
	
	try:
		if not data:
			return {"success": False, "error": "No data available"}
		
		wb = openpyxl.Workbook()
		ws = wb.active
		ws.title = "RM Report"
		
		# Styling
		header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
		header_font = Font(bold=True, color="FFFFFF", size=11)
		total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
		border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
		center_align = Alignment(horizontal='center', vertical='center')
		right_align = Alignment(horizontal='right', vertical='center')
		
		# Headers
		headers = ["Item Code", "Item Name", "UOM", "Group", "Opening", "In", "Out", "Balance", "Last Rate", "Valuation"]
		for col_num, header in enumerate(headers, 1):
			cell = ws.cell(row=1, column=col_num)
			cell.value = header
			cell.fill = header_fill
			cell.font = header_font
			cell.alignment = center_align
			cell.border = border

		current_row = 2
		for row in data:
			ws.cell(row=current_row, column=1, value=row.get('item_code')).border = border
			ws.cell(row=current_row, column=2, value=row.get('item_name')).border = border
			ws.cell(row=current_row, column=3, value=row.get('stock_uom')).border = border
			ws.cell(row=current_row, column=4, value=row.get('item_group')).border = border
			
			qty_cols = [
				row.get('opening_qty', 0), row.get('in_qty', 0), row.get('out_qty', 0), 
				row.get('balance_qty', 0), row.get('last_purchase_rate', 0), row.get('balance_value', 0)
			]
			for i, val in enumerate(qty_cols, 5):
				cell = ws.cell(row=current_row, column=i, value=float(val))
				cell.border = border
				cell.number_format = '#,##0.00'
				cell.alignment = right_align
			current_row += 1

		# Grand Total
		ws.cell(row=current_row, column=1, value="Grand Total").font = Font(bold=True)
		ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
		
		total_keys = ['opening_qty', 'in_qty', 'out_qty', 'balance_qty']
		for i, key in enumerate(total_keys, 5):
			cell = ws.cell(row=current_row, column=i, value=float(grand_total.get(key, 0)))
			cell.font = Font(bold=True)
			cell.fill = total_fill
			cell.border = border
			cell.number_format = '#,##0.00'
			cell.alignment = right_align
			
		# Skip Last Rate for grand total
		ws.cell(row=current_row, column=9).fill = total_fill
		ws.cell(row=current_row, column=9).border = border
		
		# Valuation total
		cell = ws.cell(row=current_row, column=10, value=float(grand_total.get('balance_value', 0)))
		cell.font = Font(bold=True)
		cell.fill = total_fill
		cell.border = border
		cell.number_format = '#,##0.00'
		cell.alignment = right_align
		
		# Widths
		ws.column_dimensions['A'].width = 20
		ws.column_dimensions['B'].width = 30
		for char in 'EFGHJ':
			ws.column_dimensions[char].width = 15

		file_data = BytesIO()
		wb.save(file_data)
		file_data.seek(0)
		
		filename = f"RM_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
		file_doc = frappe.get_doc({"doctype": "File", "file_name": filename, "is_private": 0, "content": file_data.getvalue()})
		file_doc.save(ignore_permissions=True)
		
		return {"success": True, "file_url": file_doc.file_url, "file_name": filename}
		
	except Exception as e:
		frappe.log_error(frappe.get_traceback(), "RM Export Error")
		return {"success": False, "error": str(e)}
